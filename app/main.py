from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Any, Callable

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .analyzer import AccountAnalyzer
from .config import BASE_DIR, settings
from .graph_client import GraphApiError
from .pdf_export import build_account_pdf, build_batch_pdf, pdf_filename_for_account

app = FastAPI(
    title="M365 Account Analyzer",
    description="Analyse tenant-only Microsoft 365 / Entra ID pour qualifier l'usage d'un compte.",
    version="1.5.0",
)

STATIC_DIR = BASE_DIR / "app" / "static"
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

_analyzer: AccountAnalyzer | None = None


def get_analyzer() -> AccountAnalyzer:
    global _analyzer
    if _analyzer is None:
        _analyzer = AccountAnalyzer()
    return _analyzer


@app.exception_handler(GraphApiError)
def graph_error_handler(_, exc: GraphApiError):
    return JSONResponse(
        status_code=exc.status_code or 500,
        content={"error": str(exc), "details": exc.details},
    )


@app.get("/", include_in_schema=False)
def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
def health() -> dict[str, Any]:
    missing = settings.validate()
    return {
        "status": "ok" if not missing else "configuration_incomplete",
        "missingEnv": missing,
        "graphPeriod": settings.graph_period,
        "reportsEnabled": settings.enable_reports,
        "cacheHours": settings.report_cache_hours,
        "exchangeEnabled": settings.enable_exchange_online,
        "exchangeAuthMode": settings.exchange_auth_mode,
        "exchangeOrganization": settings.exchange_organization,
        "exchangeCacheMinutes": settings.exchange_cache_minutes,
    }


@app.get("/api/graph-test")
def graph_test() -> dict[str, Any]:
    analyzer = get_analyzer()
    result = analyzer.graph.get(
        "/users",
        params={
            "$top": 5,
            "$select": "id,displayName,userPrincipalName,mail,accountEnabled",
        },
    )
    return {
        "status": "graph_ok",
        "users": result.get("value", []),
    }


@app.get("/api/diagnostics")
def diagnostics() -> dict[str, Any]:
    missing = settings.validate()
    base = {
        "status": "ok",
        "missingEnv": missing,
        "graphPeriod": settings.graph_period,
        "reportsEnabled": settings.enable_reports,
        "cacheHours": settings.report_cache_hours,
        "exchangeEnabled": settings.enable_exchange_online,
        "exchangeAuthMode": settings.exchange_auth_mode,
        "exchangeOrganization": settings.exchange_organization,
        "exchangeCacheMinutes": settings.exchange_cache_minutes,
        "tests": [],
    }
    if missing:
        base["status"] = "configuration_incomplete"
        base["tests"] = [
            {
                "name": "Variables .env",
                "ok": False,
                "detail": "Variables manquantes : " + ", ".join(missing),
            }
        ]
        return base

    analyzer = get_analyzer()

    def run_test(name: str, func: Callable[[], str]) -> dict[str, Any]:
        try:
            detail = func()
            return {"name": name, "ok": True, "detail": detail}
        except GraphApiError as exc:
            return {"name": name, "ok": False, "detail": str(exc), "details": exc.details}
        except Exception as exc:  # noqa: BLE001 - diagnostic admin
            return {"name": name, "ok": False, "detail": str(exc)}

    tests = [
        run_test(
            "Token Graph + lecture utilisateurs",
            lambda: f"OK, {len(analyzer.graph.get('/users', params={'$top': 3, '$select': 'id,displayName,userPrincipalName'}).get('value', []))} utilisateur(s) retournes.",
        ),
        run_test(
            "Directory.Read.All / groupes",
            lambda: f"OK, {len(analyzer.graph.get('/groups', params={'$top': 3, '$select': 'id,displayName,mail'}).get('value', []))} groupe(s) retournes.",
        ),
        run_test(
            "AuditLog.Read.All / sign-ins",
            lambda: f"OK, {len(analyzer.graph.get('/auditLogs/signIns', params={'$top': 1, '$select': 'id,createdDateTime,userPrincipalName,appDisplayName'}).get('value', []))} sign-in(s) retournes.",
        ),
        run_test(
            "Licences tenant / subscribedSkus",
            lambda: f"OK, {len(analyzer.graph.get('/subscribedSkus').get('value', []))} SKU(s) retournes.",
        ),
        run_test(
            "UserAuthenticationMethod.Read.All / MFA",
            lambda: _diagnostic_auth_methods(analyzer),
        ),
        run_test(
            "TransitiveMemberOf / roles transitifs",
            lambda: _diagnostic_transitive_member_of(analyzer),
        ),
    ]

    if settings.enable_reports:
        tests.append(
            run_test(
                "Reports.Read.All / rapport utilisateurs actifs",
                lambda: "OK, rapport lisible (" + str(len(analyzer.graph.download_report_csv("/reports/getOffice365ActiveUserDetail(period='D7')"))) + " caracteres).",
            )
        )
    else:
        tests.append(
            {
                "name": "Reports.Read.All / rapport utilisateurs actifs",
                "ok": None,
                "detail": "Test ignore : ENABLE_REPORTS=false.",
            }
        )


    exchange_diag, exchange_warnings = analyzer.exchange.diagnostic()
    if settings.enable_exchange_online:
        tests.append(
            {
                "name": "Exchange Online API / mailbox lookup",
                "ok": bool(exchange_diag.get("ok")),
                "detail": exchange_diag.get("detail") or exchange_diag.get("error") or str(exchange_diag),
                "details": exchange_diag,
            }
        )
        for warning in exchange_warnings:
            tests.append({"name": "Exchange Online avertissement", "ok": None, "detail": warning})
    else:
        tests.append(
            {
                "name": "Exchange Online API / mailbox lookup",
                "ok": None,
                "detail": "Test ignore : ENABLE_EXCHANGE_ONLINE=false.",
            }
        )

    base["tests"] = tests
    base["status"] = "ok" if all(test.get("ok") in {True, None} for test in tests) else "partial"
    return base


@app.get("/api/account")
def analyze_account(identifier: str = Query(..., description="UPN, mail, id objet, samAccountName local synchronise, alias ou nom affiche")):
    if not identifier.strip():
        raise HTTPException(status_code=400, detail="Identifiant obligatoire.")
    analyzer = get_analyzer()
    return analyzer.analyze(identifier.strip())


@app.get("/api/account/pdf")
def analyze_account_pdf(identifier: str = Query(..., description="UPN, mail, id objet, samAccountName local synchronise, alias ou nom affiche")) -> Response:
    if not identifier.strip():
        raise HTTPException(status_code=400, detail="Identifiant obligatoire.")
    analyzer = get_analyzer()
    result = analyzer.analyze(identifier.strip())
    pdf = build_account_pdf(result)
    filename = pdf_filename_for_account(result, identifier)
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/batch")
async def analyze_batch(file: UploadFile = File(...)):
    raw = await file.read()
    text = raw.decode("utf-8-sig", errors="replace")
    identifiers = _extract_identifiers_from_csv_or_text(text)
    if not identifiers:
        raise HTTPException(status_code=400, detail="Aucun identifiant trouve dans le fichier.")

    analyzer = get_analyzer()
    results = []
    for identifier in identifiers[:200]:
        try:
            results.append({"identifier": identifier, "ok": True, "result": analyzer.analyze(identifier)})
        except Exception as exc:  # noqa: BLE001 - retour utile en analyse par lot
            results.append({"identifier": identifier, "ok": False, "error": str(exc)})

    return {"count": len(results), "summary": _batch_summary(results), "results": results}


@app.post("/api/batch/export")
async def analyze_batch_export(file: UploadFile = File(...)):
    raw = await file.read()
    text = raw.decode("utf-8-sig", errors="replace")
    identifiers = _extract_identifiers_from_csv_or_text(text)
    if not identifiers:
        raise HTTPException(status_code=400, detail="Aucun identifiant trouve dans le fichier.")

    analyzer = get_analyzer()
    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=[
            "identifier_recherche",
            "ok",
            "nom",
            "upn",
            "mail",
            "compte_active",
            "sync_ad",
            "samaccountname",
            "derniere_synchro_ad",
            "statut",
            "score",
            "confiance",
            "risque",
            "dernier_signal",
            "dernier_signal_jours",
            "licences_count",
            "exchange_type",
            "exchange_shared",
            "exchange_last_user_action",
            "exchange_last_logon",
            "exchange_total_size",
            "exchange_fullaccess_count",
            "exchange_sendas_count",
            "exchange_sendonbehalf_count",
            "groupes_count",
            "roles_count",
            "technique_probable",
            "critique_potentiel",
            "recommandation",
            "avertissements",
            "erreur",
        ],
        delimiter=";",
    )
    writer.writeheader()

    for identifier in identifiers[:200]:
        try:
            result = analyzer.analyze(identifier)
            identity = result.get("identity", {})
            onprem = identity.get("onPremises", {})
            conclusion = result.get("conclusion", {})
            licenses = result.get("licenses", {})
            exchange = result.get("exchange", {})
            memberships = result.get("memberships", {})
            latest = conclusion.get("latestSignal") or {}
            writer.writerow(
                {
                    "identifier_recherche": identifier,
                    "ok": "oui",
                    "nom": identity.get("displayName"),
                    "upn": identity.get("userPrincipalName"),
                    "mail": identity.get("mail"),
                    "compte_active": identity.get("accountEnabled"),
                    "sync_ad": onprem.get("syncEnabled"),
                    "samaccountname": onprem.get("samAccountName"),
                    "derniere_synchro_ad": onprem.get("lastSyncDateTime"),
                    "statut": conclusion.get("status"),
                    "score": conclusion.get("score"),
                    "confiance": conclusion.get("confidence"),
                    "risque": conclusion.get("riskLevel"),
                    "dernier_signal": latest.get("date"),
                    "dernier_signal_jours": conclusion.get("latestSignalDaysAgo"),
                    "licences_count": licenses.get("count"),
                    "exchange_type": exchange.get("recipientTypeDetails"),
                    "exchange_shared": exchange.get("isSharedMailbox"),
                    "exchange_last_user_action": exchange.get("lastUserActionTime"),
                    "exchange_last_logon": exchange.get("lastLogonTime"),
                    "exchange_total_size": exchange.get("totalItemSize"),
                    "exchange_fullaccess_count": exchange.get("fullAccessCount"),
                    "exchange_sendas_count": exchange.get("sendAsCount"),
                    "exchange_sendonbehalf_count": exchange.get("sendOnBehalfCount"),
                    "groupes_count": memberships.get("groupsCount"),
                    "roles_count": memberships.get("directoryRolesCount"),
                    "technique_probable": conclusion.get("isTechnicalCandidate"),
                    "critique_potentiel": conclusion.get("isCriticalCandidate"),
                    "recommandation": conclusion.get("recommendation"),
                    "avertissements": " | ".join(result.get("warnings") or []),
                    "erreur": "",
                }
            )
        except Exception as exc:  # noqa: BLE001
            writer.writerow(
                {
                    "identifier_recherche": identifier,
                    "ok": "non",
                    "nom": "",
                    "upn": "",
                    "mail": "",
                    "compte_active": "",
                    "sync_ad": "",
                    "samaccountname": "",
                    "derniere_synchro_ad": "",
                    "statut": "",
                    "score": "",
                    "confiance": "",
                    "risque": "",
                    "dernier_signal": "",
                    "dernier_signal_jours": "",
                    "licences_count": "",
                    "exchange_type": "",
                    "exchange_shared": "",
                    "exchange_last_user_action": "",
                    "exchange_last_logon": "",
                    "exchange_total_size": "",
                    "exchange_fullaccess_count": "",
                    "exchange_sendas_count": "",
                    "exchange_sendonbehalf_count": "",
                    "groupes_count": "",
                    "roles_count": "",
                    "technique_probable": "",
                    "critique_potentiel": "",
                    "recommandation": "",
                    "avertissements": "",
                    "erreur": str(exc),
                }
            )

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=m365_account_analyzer_export.csv"},
    )


@app.post("/api/batch/export-pdf")
async def analyze_batch_export_pdf(file: UploadFile = File(...)):
    raw = await file.read()
    text = raw.decode("utf-8-sig", errors="replace")
    identifiers = _extract_identifiers_from_csv_or_text(text)
    if not identifiers:
        raise HTTPException(status_code=400, detail="Aucun identifiant trouve dans le fichier.")

    analyzer = get_analyzer()
    results = []
    for identifier in identifiers[:200]:
        try:
            results.append({"identifier": identifier, "ok": True, "result": analyzer.analyze(identifier)})
        except Exception as exc:  # noqa: BLE001 - le PDF doit rester exportable meme si une ligne echoue
            results.append({"identifier": identifier, "ok": False, "error": str(exc)})

    pdf = build_batch_pdf(results, summary=_batch_summary(results))
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="m365_account_analyzer_export.pdf"'},
    )


@app.post("/api/batch/export-xlsx")
async def analyze_batch_export_xlsx(file: UploadFile = File(...)):
    raw = await file.read()
    text = raw.decode("utf-8-sig", errors="replace")
    identifiers = _extract_identifiers_from_csv_or_text(text)
    if not identifiers:
        raise HTTPException(status_code=400, detail="Aucun identifiant trouve dans le fichier.")

    analyzer = get_analyzer()
    results = []
    for identifier in identifiers[:200]:
        try:
            results.append({"identifier": identifier, "ok": True, "result": analyzer.analyze(identifier)})
        except Exception as exc:  # noqa: BLE001
            results.append({"identifier": identifier, "ok": False, "error": str(exc)})

    rows = [_flatten_batch_item(item) for item in results]
    summary = _batch_summary(results)

    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    ws.append(["Indicateur", "Valeur"])
    for key, value in summary.items():
        ws.append([key, value])
    _format_sheet(ws)

    ws2 = wb.create_sheet("Détails")
    headers = list(rows[0].keys()) if rows else ["identifier_recherche"]
    ws2.append(headers)
    for row in rows:
        ws2.append([row.get(h, "") for h in headers])
    _format_sheet(ws2)

    ws3 = wb.create_sheet("Critiques")
    ws3.append(headers)
    for row in rows:
        if str(row.get("critique_potentiel", "")).lower() == "true" or str(row.get("privilegie", "")).lower() == "true" or "critique" in str(row.get("statut", "")).lower():
            ws3.append([row.get(h, "") for h in headers])
    _format_sheet(ws3)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=m365_account_analyzer_export.xlsx"},
    )


def _diagnostic_auth_methods(analyzer: AccountAnalyzer) -> str:
    users = analyzer.graph.get("/users", params={"$top": 1, "$select": "id,userPrincipalName"}).get("value", [])
    if not users:
        return "Aucun utilisateur disponible pour tester les methodes MFA."
    user = users[0]
    result = analyzer.graph.get(f"/users/{analyzer.graph.quote_id(user['id'])}/authentication/methods")
    return f"OK, {len(result.get('value', []))} methode(s) retournee(s) pour {user.get('userPrincipalName')}."


def _diagnostic_transitive_member_of(analyzer: AccountAnalyzer) -> str:
    users = analyzer.graph.get("/users", params={"$top": 1, "$select": "id,userPrincipalName"}).get("value", [])
    if not users:
        return "Aucun utilisateur disponible pour tester transitiveMemberOf."
    user = users[0]
    result = analyzer.graph.get(f"/users/{analyzer.graph.quote_id(user['id'])}/transitiveMemberOf", params={"$top": 1})
    return f"OK, {len(result.get('value', []))} objet(s) retourne(s) pour {user.get('userPrincipalName')}."


def _flatten_batch_item(item: dict[str, Any]) -> dict[str, Any]:
    if not item.get("ok"):
        return {
            "identifier_recherche": item.get("identifier"),
            "ok": "non",
            "erreur": item.get("error"),
        }
    result = item.get("result") or item
    identity = result.get("identity", {})
    onprem = identity.get("onPremises", {})
    conclusion = result.get("conclusion", {})
    licenses = result.get("licenses", {})
    memberships = result.get("memberships", {})
    security = result.get("security", {})
    ownership = result.get("ownership", {})
    exchange = result.get("exchange", {})
    latest = conclusion.get("latestSignal") or {}
    return {
        "identifier_recherche": item.get("identifier") or result.get("searchedIdentifier"),
        "ok": "oui",
        "nom": identity.get("displayName"),
        "upn": identity.get("userPrincipalName"),
        "mail": identity.get("mail"),
        "compte_active": identity.get("accountEnabled"),
        "sync_ad": onprem.get("syncEnabled"),
        "samaccountname": onprem.get("samAccountName"),
        "derniere_synchro_ad": onprem.get("lastSyncDateTime"),
        "statut": conclusion.get("status"),
        "score": conclusion.get("score"),
        "confiance": conclusion.get("confidence"),
        "risque": conclusion.get("riskLevel"),
        "dernier_signal": latest.get("date"),
        "dernier_signal_jours": conclusion.get("latestSignalDaysAgo"),
        "licences_count": licenses.get("count"),
        "exchange_trouve": exchange.get("found") or exchange.get("recipientFound"),
        "exchange_type": exchange.get("recipientTypeDetails"),
        "exchange_boite": exchange.get("hasMailbox"),
        "exchange_boite_partagee": exchange.get("isSharedMailbox"),
        "exchange_smtp_principal": exchange.get("primarySmtpAddress"),
        "exchange_derniere_action": exchange.get("lastUserActionTime"),
        "exchange_dernier_logon": exchange.get("lastLogonTime"),
        "exchange_taille": exchange.get("totalItemSize"),
        "exchange_elements": exchange.get("itemCount"),
        "exchange_fullaccess_count": exchange.get("fullAccessCount"),
        "exchange_sendas_count": exchange.get("sendAsCount"),
        "exchange_sendonbehalf_count": exchange.get("sendOnBehalfCount"),
        "exchange_delegations": exchange.get("delegationsFound"),
        "groupes_directs_count": memberships.get("groupsCount"),
        "roles_directs_count": memberships.get("directoryRolesCount"),
        "roles_transitifs_count": memberships.get("transitiveDirectoryRolesCount"),
        "privilegie": conclusion.get("isPrivileged"),
        "mfa_detecte": conclusion.get("hasMfaMethod"),
        "methodes_mfa": ", ".join(security.get("strongMethods") or []),
        "technique_probable": conclusion.get("isTechnicalCandidate"),
        "critique_potentiel": conclusion.get("isCriticalCandidate"),
        "proprietaire_probable": ownership.get("probableOwner"),
        "manager": (ownership.get("manager") or {}).get("displayName") if ownership.get("manager") else "",
        "objets_possedes_count": ownership.get("ownedObjectsCount"),
        "recommandation": conclusion.get("recommendation"),
        "avertissements": " | ".join(result.get("warnings") or []),
        "erreur": "",
    }


def _batch_summary(results: list[dict[str, Any]]) -> dict[str, Any]:
    rows = [_flatten_batch_item(item) for item in results]
    ok_rows = [r for r in rows if r.get("ok") == "oui"]
    def count_where(predicate):
        return sum(1 for r in ok_rows if predicate(r))
    return {
        "analyses": len(results),
        "ok": len(ok_rows),
        "erreurs": len(results) - len(ok_rows),
        "actifs_recents": count_where(lambda r: "actif" in str(r.get("statut", "")).lower()),
        "dormants_m365": count_where(lambda r: "dormant" in str(r.get("statut", "")).lower()),
        "techniques": count_where(lambda r: str(r.get("technique_probable", "")).lower() == "true"),
        "critiques_infra": count_where(lambda r: str(r.get("critique_potentiel", "")).lower() == "true" or "critique" in str(r.get("statut", "")).lower()),
        "privilegies": count_where(lambda r: str(r.get("privilegie", "")).lower() == "true"),
        "mfa_detecte": count_where(lambda r: str(r.get("mfa_detecte", "")).lower() == "true"),
        "sans_mfa_lisible": count_where(lambda r: str(r.get("mfa_detecte", "")).lower() != "true"),
        "boites_exchange": count_where(lambda r: str(r.get("exchange_boite", "")).lower() == "true"),
        "boites_partagees": count_where(lambda r: str(r.get("exchange_boite_partagee", "")).lower() == "true"),
        "delegations_exchange": count_where(lambda r: str(r.get("exchange_delegations", "")).lower() == "true"),
    }


def _format_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _extract_identifiers_from_csv_or_text(text: str) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return []

    sample = "\n".join(lines[:5])
    dialect = csv.excel
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        pass

    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    identifiers: list[str] = []

    if reader.fieldnames:
        normalized = {name.lower().strip(): name for name in reader.fieldnames if name}
        preferred_columns = [
            "identifiant (upn)",
            "upn",
            "userprincipalname",
            "user principal name",
            "adresse email",
            "adresse mail",
            "email",
            "mail",
            "samaccountname",
            "nom complet",
            "displayname",
            "nom",
        ]
        selected_column = None
        for key in preferred_columns:
            if key in normalized:
                selected_column = normalized[key]
                break
        if selected_column:
            for row in reader:
                value = (row.get(selected_column) or "").strip()
                if value:
                    identifiers.append(value)
            return _dedupe(identifiers)

    # Fallback : une ligne = un identifiant. Utile pour un simple copier-coller.
    return _dedupe(lines)


def _dedupe(values: list[str]) -> list[str]:
    seen = set()
    result = []
    for value in values:
        key = value.lower()
        if key not in seen:
            seen.add(key)
            result.append(value)
    return result
